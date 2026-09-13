"""Generate the non-demo upload fixtures used by scripts/browser_uploads.py.

Run inside the backend image so the pinned pypdf/python-docx versions produce the files.
The runtime image ships only the entrypoint, so mount this script and an output directory:

    mkdir -p .fixtures
    docker compose run --rm --no-deps --user root \
        -v "$PWD/scripts/make_upload_fixtures.py":/tmp/make_upload_fixtures.py \
        -v "$PWD/.fixtures":/out backend python /tmp/make_upload_fixtures.py --out /out

The output is disposable test input, not source material; keep it out of Git.
"""
import argparse
import io
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

SOFTWARE_TEXT = "The Python software package contains a script and a repository for a local demonstration of reuse."


def text_pdf(text: str) -> bytes:
    """A minimal, genuinely text-based single-page PDF (no OCR, no image layer)."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def method_docx() -> bytes:
    document = Document()
    document.add_heading("Fictional protocol handbook", 0)
    document.add_paragraph("This method describes a reusable workflow and procedure for reviewing archived field notes before reuse.")
    document.add_paragraph("A separate lecture and assignment turn the same protocol into a teaching tutorial for new staff.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Evidence"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    out = parser.parse_args().out
    out.mkdir(parents=True, exist_ok=True)
    (out / "upload-software.pdf").write_bytes(text_pdf(SOFTWARE_TEXT))
    (out / "upload-method.docx").write_bytes(method_docx())
    # Fictional prose with no rule term, so a run must report no candidates rather than invent one.
    (out / "irrelevant.txt").write_text(
        "Yesterday the weather stayed mild and the kitchen clock needed a new battery.\n\n"
        "A neighbour repainted a fence, then everyone went home for supper.\n", encoding="utf-8")
    # Truncated/By-name-only binaries: the parser must reject these with a readable message.
    (out / "malformed.pdf").write_bytes(b"%PDF-1.7\nthis is not really a pdf body\n%%EOF\n")
    (out / "malformed.docx").write_bytes(b"PK\x03\x04 not actually a docx container at all")
    for path in sorted(out.iterdir()):
        print(f"{path.name} {path.stat().st_size}")


if __name__ == "__main__":
    main()
