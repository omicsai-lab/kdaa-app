import io
from uuid import uuid4
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from kdaa.errors import KDAAError
from kdaa.parsing import MAX_BYTES, MAX_CHARACTERS, normalize, parse_document
from kdaa.providers.local import LocalAssetStore

def pdf_fixture(text="The Python package contains a script and an example for a local software demonstration."):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output); return output.getvalue()

def docx_fixture():
    doc = Document(); doc.add_heading("Fictional teaching source", 0)
    doc.add_paragraph("A lecture and assignment explain how to inspect exact source quotations in teaching material.")
    table = doc.add_table(rows=1, cols=2); table.cell(0,0).text = "Topic"; table.cell(0,1).text = "Evidence"
    out = io.BytesIO(); doc.save(out); return out.getvalue()

@pytest.mark.parametrize("name", ["input.txt", "input.md", "INPUT.TXT", "notes.MD"])
def test_utf8_text(name):
    result = parse_document(name, b"\xef\xbb\xbfHello\r\n\r\n  A  method.  \r\n")
    assert result.text == "Hello\n\nA method."

@pytest.mark.parametrize("name", ["../secret.txt", "/tmp/a.md", "a\\b.txt", "bad\x00.txt", "", "test.csv", "a.exe"])
def test_bad_filenames(name):
    with pytest.raises(KDAAError): parse_document(name, b"Some content")

@pytest.mark.parametrize("content", [b"", b"    \r\n", b"\xff\xfe", b"abc\x00xyz", b"abc\x01xyz"])
def test_invalid_text(content):
    with pytest.raises(KDAAError): parse_document("a.txt", content)

def test_limits():
    with pytest.raises(KDAAError): parse_document("a.txt", b"a" * (MAX_BYTES+1))
    with pytest.raises(KDAAError): normalize("x" * (MAX_CHARACTERS+1))

def test_nfc_and_whitespace():
    assert normalize("Cafe\u0301\r\n\r\n\r\n  next\tline  ") == "Café\n\nnext line"

def test_real_pdf_extraction():
    result = parse_document("text.pdf", pdf_fixture())
    assert "Python package" in result.text
    assert "no OCR" in result.warnings[0]

def test_blank_pdf_is_not_successful_import():
    w = PdfWriter(); w.add_blank_page(width=100, height=100); b = io.BytesIO(); w.write(b)
    with pytest.raises(KDAAError, match="No readable text"):
        parse_document("scan.pdf", b.getvalue())

def test_encrypted_pdf_rejected():
    w = PdfWriter(); w.add_blank_page(width=100, height=100); w.encrypt("secret"); b = io.BytesIO(); w.write(b)
    with pytest.raises(KDAAError, match="Encrypted"):
        parse_document("secret.pdf", b.getvalue())

def test_real_docx_extraction():
    result = parse_document("input.docx", docx_fixture())
    assert "lecture and assignment" in result.text
    assert "Topic | Evidence" in result.text

@pytest.mark.parametrize("name", ["a.pdf", "a.docx"])
def test_format_content_mismatch(name):
    with pytest.raises(KDAAError): parse_document(name, b"not a binary document")

def test_local_storage_safe_keys_and_immutability(tmp_path):
    store = LocalAssetStore(tmp_path)
    key = f"{uuid4()}/{uuid4()}.raw"
    store.put(key, b"original")
    assert store.read(key) == b"original"
    with pytest.raises(KDAAError): store.put(key, b"changed")
    with pytest.raises(KDAAError): store.read("../secret")
    assert store.ready()
    store.delete(key)
    with pytest.raises(KDAAError): store.read(key)
